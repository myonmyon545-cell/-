#!/usr/bin/env python3
"""Read-only extraction for the T11 E20-02 summary sheet."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

def sha(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):
            h.update(chunk)
    return h.hexdigest()

def cell(wb, sheet, ref):
    return wb[sheet][ref].value

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--source-copy',type=Path,required=True)
    ap.add_argument('--first-round',type=Path,required=True)
    ap.add_argument('--out',type=Path,required=True)
    a=ap.parse_args()
    s01=load_workbook(a.source_copy,read_only=True,data_only=True)
    s02=load_workbook(a.first_round,read_only=True,data_only=True)
    required=['表紙・目次','データの見方','Ⅰ基礎情報','Ⅱ調査項目データ（第１群）',
      'Ⅱ調査項目データ（第２群）','Ⅱ調査項目データ（第３群）','Ⅱ調査項目データ（第４群）',
      'Ⅱ調査項目データ（第５群）','Ⅱ調査項目データ（その他）','Ⅱ（３）５項目',
      'Ⅲ審査判定データ（一次・二次判定）','Ⅲ審査判定データ（変更率）','Ⅲ審査判定データ（有効期間）','Ⅳ事務データ']
    if s01.sheetnames[1:] != required:
        raise ValueError('original 14 worksheet order/name mismatch')
    if 'Ⅰ基礎情報' not in s02.sheetnames or 'Ⅲ審査判定データ（一次・二次判定）' not in s02.sheetnames:
        raise ValueError('first-round source structure mismatch')
    i=s01['Ⅰ基礎情報']; i2=s02['Ⅰ基礎情報']
    values=[
      {'id':'R02-AGE65','label':'65歳以上人口','category':'人口構成','unit':'人','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','G130'),'pref':None,'national':None,'prior':None,'ref':'S01 Ⅰ基礎情報!G130','comparison':'人口実数。県・全国比較は掲載しない。提供期間差対象外。'},
      {'id':'R02-AGING','label':'高齢化率（65歳以上/総人口）','category':'人口構成','unit':'比率','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','T130'),'pref':cell(s01,'Ⅰ基礎情報','W130'),'national':cell(s01,'Ⅰ基礎情報','Z130'),'prior':None,'ref':'S01 Ⅰ基礎情報!T130/W130/Z130','comparison':'同一基準日の地域構成比較。県差・全国差はpp。提供期間差対象外。'},
      {'id':'R02-AGE75','label':'75歳以上割合（65歳以上人口中）','category':'人口構成','unit':'比率','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','K131'),'pref':cell(s01,'Ⅰ基礎情報','N131'),'national':cell(s01,'Ⅰ基礎情報','Q131'),'prior':None,'ref':'S01 Ⅰ基礎情報!K131/N131/Q131','comparison':'同一基準日の地域構成比較。県差・全国差はpp。提供期間差対象外。'},
      {'id':'R02-AGE85','label':'85歳以上割合（65歳以上人口中）','category':'人口構成','unit':'比率','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','K132'),'pref':cell(s01,'Ⅰ基礎情報','N132'),'national':cell(s01,'Ⅰ基礎情報','Q132'),'prior':None,'ref':'S01 Ⅰ基礎情報!K132/N132/Q132','comparison':'同一基準日の地域構成比較。県差・全国差はpp。提供期間差対象外。'},
      {'id':'R03-CERT','label':'認定率（第1号/65歳以上）','category':'認定状況','unit':'比率','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','U261'),'pref':cell(s01,'Ⅰ基礎情報','X261'),'national':cell(s01,'Ⅰ基礎情報','AA261'),'prior':None,'ref':'S01 Ⅰ基礎情報!U261/X261/AA261','comparison':'同一基準日の地域比較。差はpp。S02の同基準値を期間差として比較しない。'},
      {'id':'R03-LIGHT','label':'軽度認定率（第1号/65歳以上）','category':'認定状況','unit':'比率','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','U262'),'pref':cell(s01,'Ⅰ基礎情報','X262'),'national':cell(s01,'Ⅰ基礎情報','AA262'),'prior':None,'ref':'S01 Ⅰ基礎情報!U262/X262/AA262','comparison':'同一基準日の地域比較。差はpp。'},
      {'id':'R03-MODSEV','label':'中・重度認定率（第1号/65歳以上）','category':'認定状況','unit':'比率','period':'基準日 2025-01-01','kind':'通常','city':cell(s01,'Ⅰ基礎情報','U263'),'pref':cell(s01,'Ⅰ基礎情報','X263'),'national':cell(s01,'Ⅰ基礎情報','AA263'),'prior':None,'ref':'S01 Ⅰ基礎情報!U263/X263/AA263','comparison':'同一基準日の地域比較。差はpp。'},
      {'id':'R06-PRIMARY-N','label':'一次判定対象件数（C03既存反映）','category':'審査判定','unit':'件','period':'S01 2025-04-01～2025-09-30 / S02 2024-10-01～2025-03-31','kind':'通常','city':cell(s01,'Ⅲ審査判定データ（一次・二次判定）','G136'),'pref':None,'national':None,'prior':cell(s02,'Ⅲ審査判定データ（一次・二次判定）','G136'),'ref':'S01 Ⅲ審査判定データ（一次・二次判定）!G136; 構成計 S01 Ⅰ!G513:G520','prior_ref':'S02 Ⅲ審査判定データ（一次・二次判定）!G136','comparison':'同じ一次判定対象件数の提供期間差。C03は既存作業コピーに適用済み、再適用なし。'}
    ]
    cells=[cell(s01,'Ⅰ基礎情報',f'G{r}') for r in range(513,522)]
    sum8=sum(cells[:8])
    total=cell(s01,'Ⅲ審査判定データ（一次・二次判定）','G136')
    if sum8 != total or cells[8] != total or values[-1]['city'] != total:
        raise ValueError(f'C03_existing_reconciliation_failed: sum8={sum8}, row9={cells[8]}, total={total}')
    values[-1]['prior_diff']=values[-1]['city']-values[-1]['prior']
    values[-1]['c03_cells']=cells
    result={
      'schema':'care_t11_e20_02_input/v1',
      'source_copy_sha256':sha(a.source_copy),
      'source_copy_sheets':len(s01.sheetnames),
      'original_sheet_names':s01.sheetnames[1:],
      'first_round_sha256':sha(a.first_round),
      's01_input_id':'libfile_2f275b2a53d481919fb7d942fbecc486',
      's01_input_sha256':'5721ee2b38cb8f55f973f149fe32f67dcfe5453645fcfdccdb934b169bcd2a2d',
      's02_input_id':'libfile_49bf2650277881919afc2e41b9afaca2',
      's02_input_sha256':'9b29b315bf310770295e8fe60ba2953362f3cff21b3d78adfe255c42f2c26304',
      'periods':{'S01':['2025-04-01','2025-09-30'],'S02':['2024-10-01','2025-03-31'],'population_reference_date':'2025-01-01','warehouse_cutoff':'2025-11-30'},
      'records':values,
      'c03_reconciliation':{'sum_I_G513_G520':sum8,'I_G521':cells[8],'III_G136':total,'result':'PASS','reapplied':False}
    }
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'records':len(values),'c03_reconciliation':result['c03_reconciliation'],'sha256':sha(a.source_copy)},ensure_ascii=False))
if __name__=='__main__': main()

